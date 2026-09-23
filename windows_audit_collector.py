#!/usr/bin/env python3
"""
=============================================================================
ISO 9001 / ITGC Technical Audit - Stage 1 Reconnaissance & Telemetry Intake
=============================================================================
Comprehensive passive evidence collector for Windows operating environments.
Designed for the RECONNAISSANCE STAGE (Preliminary Intake):
  - Performs 100% LIVE system telemetry extraction on Windows.
  - Does NOT finalize compliance pass/fail determinations.
  - Gathers 20 comprehensive audit control points across Hardware, OS,
    Network, Storage, Access Control, Encryption, Logging, and Patch Health.
  - Formats all observations with Review Team Action Guides.
  - Generates reports in .txt, .docx, and .pdf, bundled into a final .zip.
=============================================================================
"""

import argparse
import datetime
import json
import os
import platform
import re
import subprocess
import sys
import zipfile

# Document formatting libraries
try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
        KeepTogether
    )
    from reportlab.pdfgen import canvas
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False


# ============================================================================
# 20 AUDIT CONTROL DEFINITIONS (ISO 9001 / ITGC / CIS MAPPINGS)
# ============================================================================
AUDIT_ITEMS_DEF = [
    # --- 12 Core Items from Sheet ---
    {
        "id": "HW_SYSTEM",
        "clause": "7.1.3 (Infrastructure)",
        "category": "Hardware Baseline",
        "data_point": "System Hostname, Manufacturer, & Model",
        "command": "Get-CimInstance Win32_ComputerSystem | Select Name, Manufacturer, Model, TotalPhysicalMemory, Domain",
        "reviewer_action": "Cross-reference hostname and model against physical asset tag register and CMDB records.",
    },
    {
        "id": "HW_DISK_RAM",
        "clause": "7.1.3 (Infrastructure)",
        "category": "Hardware Baseline",
        "data_point": "Total RAM & Disk Capacity / Health",
        "command": "Get-CimInstance Win32_LogicalDisk | Select DeviceId, DriveType, Size, FreeSpace, FileSystem, VolumeName",
        "reviewer_action": "Verify system drive headroom (>15% free) to prevent logging bottlenecks or system crashes.",
    },
    {
        "id": "OS_VERSION",
        "clause": "7.1.3 & 8.5.1",
        "category": "OS Integrity",
        "data_point": "Windows Version & Build Number",
        "command": "Get-CimInstance Win32_OperatingSystem | Select Caption, Version, BuildNumber, OSArchitecture, InstallDate",
        "reviewer_action": "Confirm Windows build is actively supported by vendor lifecycle policy and no legacy OS is active.",
    },
    {
        "id": "OS_HOTFIX",
        "clause": "7.1.3 & 8.5.1",
        "category": "OS Integrity",
        "data_point": "Last Quality Update Installation Date",
        "command": "Get-HotFix | Sort InstalledOn -Descending | Select -First 10 HotFixID, Description, InstalledOn, InstalledBy",
        "reviewer_action": "Verify recent cumulative update installation against corporate 30-day patch management SLA.",
    },
    {
        "id": "SEC_ANTIVIRUS",
        "clause": "6.1 (Risk Management)",
        "category": "Antivirus Security",
        "data_point": "AV Engine & Definition Status",
        "command": "Get-CimInstance -Namespace root/SecurityCenter2 -ClassName AntivirusProduct; Get-MpComputerStatus",
        "reviewer_action": "Verify Real-Time Protection is enabled and signature definitions updated within the last 24 hours.",
    },
    {
        "id": "SEC_FIREWALL",
        "clause": "6.1 (Risk Management)",
        "category": "Network Security",
        "data_point": "Windows Firewall Profile State",
        "command": "Get-NetFirewallProfile | Select Name, Enabled, DefaultInboundAction, DefaultOutboundAction",
        "reviewer_action": "Ensure Domain, Private, and Public profiles are active with Public profile blocking unsolicited inbound traffic.",
    },
    {
        "id": "ACC_LOCAL_USERS",
        "clause": "7.5.3 (Doc Info)",
        "category": "Access Control",
        "data_point": "Local Users & Enabled Status",
        "command": "Get-LocalUser | Select Name, Enabled, LastLogon, Description, PasswordRequired",
        "reviewer_action": "Review all active accounts; verify generic or shared accounts are disabled or tied to named personnel.",
    },
    {
        "id": "ACC_ADMIN_MEMBERS",
        "clause": "7.5.3 (Doc Info)",
        "category": "Access Control",
        "data_point": "Local Administrator Group Members",
        "command": "Get-LocalGroupMember -Group 'Administrators' | Select Name, PrincipalSource, ObjectClass",
        "reviewer_action": "Validate administrative accounts against least privilege and confirm individual job role authorization.",
    },
    {
        "id": "STO_SMB_SHARES",
        "clause": "7.5.3 (Doc Info)",
        "category": "Data Storage",
        "data_point": "Mapped Network Drives & File Shares",
        "command": "Get-SmbMapping | Select LocalPath, RemotePath, Status; Get-SmbShare | Select Name, Path",
        "reviewer_action": "Confirm mapped network drives connect securely to approved central storage/NAS for master ISO documentation.",
    },
    {
        "id": "NET_IP_CONFIG",
        "clause": "7.1.3 & 7.1.4",
        "category": "Network Devices",
        "data_point": "IP Configuration & Subnet Routing",
        "command": "Get-NetIPAddress -AddressFamily IPv4 | Where IPAddress -notlike '127.*'; Get-NetRoute (0.0.0.0/0)",
        "reviewer_action": "Cross-check assigned IP address, subnet mask, and default gateway against documented network topology.",
    },
    {
        "id": "PRN_QUEUES",
        "clause": "7.1.3 & 7.1.4",
        "category": "Printers / Peripherals",
        "data_point": "Connected Printers & Print Queues",
        "command": "Get-Printer | Select Name, DriverName, PortName, Shared, PrinterStatus",
        "reviewer_action": "Verify connected printers route through central print servers for audited print logging.",
    },
    {
        "id": "NET_INFRA_BACKUP",
        "clause": "7.1.3 (Infrastructure)",
        "category": "Network Infrastructure",
        "data_point": "Router / Firewall Configurations & Drop Zone Backups",
        "command": "Drop Zone repository inspection (Search for .conf / .cfg gateway backup archives)",
        "reviewer_action": "Confirm existence and timestamp of recent gateway/firewall configuration backups on the drop zone.",
    },

    # --- 8 Additional Enterprise Audit Control Points ---
    {
        "id": "SEC_BITLOCKER",
        "clause": "6.1 & 7.1.3",
        "category": "Data Protection",
        "data_point": "BitLocker Drive Encryption Status",
        "command": "Get-BitLockerVolume | Select MountPoint, VolumeStatus, ProtectionStatus, EncryptionMethod",
        "reviewer_action": "Confirm system drive (C:) and fixed data drives have BitLocker protection enabled to prevent physical data compromise.",
    },
    {
        "id": "OS_TIME_SYNC",
        "clause": "7.1.5 (Traceability)",
        "category": "Measurement & Traceability",
        "data_point": "NTP Time Synchronization Status",
        "command": "w32tm /query /status | Select-String 'Source|Stratum|Leap Indicator'",
        "reviewer_action": "Verify host clock is synchronized with authoritative corporate or NTP time source to ensure log audit integrity.",
    },
    {
        "id": "SEC_UAC",
        "clause": "6.1 & 7.5.3",
        "category": "Security Hardening",
        "data_point": "User Account Control (UAC) Elevation Policy",
        "command": "Get-ItemProperty HKLM:\\Software\\Microsoft\\Windows\\CurrentVersion\\Policies\\System -Name EnableLUA",
        "reviewer_action": "Verify EnableLUA=1 to ensure User Account Control enforces permission elevation prompts against unauthorized execution.",
    },
    {
        "id": "OS_UPDATE_SERVICE",
        "clause": "8.5.1 & 7.1.3",
        "category": "Maintenance Operations",
        "data_point": "Windows Update Service (wuauserv) Health",
        "command": "Get-Service wuauserv | Select Name, DisplayName, Status, StartType",
        "reviewer_action": "Ensure Windows Update service (wuauserv) is not disabled or stopped, supporting continual maintenance operations.",
    },
    {
        "id": "ACC_PASSWORD_POLICY",
        "clause": "7.5.3 (Doc Info)",
        "category": "Access Control",
        "data_point": "Local Password Policy & Lockout Threshold",
        "command": "net accounts",
        "reviewer_action": "Review minimum password length, maximum password age, and lockout threshold to prevent credential brute forcing.",
    },
    {
        "id": "NET_LISTENING_PORTS",
        "clause": "6.1 & 7.1.4",
        "category": "Network Security",
        "data_point": "Active Listening Network Ports & Exposed Services",
        "command": "Get-NetTCPConnection -State Listen | Select LocalAddress, LocalPort, OwningProcess | Group LocalPort",
        "reviewer_action": "Inspect listening TCP ports to verify only authorized network services are listening externally.",
    },
    {
        "id": "SEC_RDP_CONFIG",
        "clause": "6.1 & 7.5.3",
        "category": "Remote Access",
        "data_point": "Remote Desktop (RDP) & NLA Status",
        "command": "Get-ItemProperty 'HKLM:\\System\\CurrentControlSet\\Control\\Terminal Server' -Name fDenyTSConnections",
        "reviewer_action": "Verify RDP policy complies with organization standards and Network Level Authentication (NLA) is strictly required.",
    },
    {
        "id": "LOG_AUDIT_HEALTH",
        "clause": "7.5.3 & 9.1",
        "category": "Monitoring & Measurement",
        "data_point": "Security & System Event Log Retention Health",
        "command": "Get-WinEvent -ListLog Security, System | Select LogName, RecordCount, IsLogFull, MaximumSizeInBytes",
        "reviewer_action": "Verify Windows Security and System event logs are recording records and have not reached full capacity.",
    },
]


# ============================================================================
# ROBUST POWERSHELL EXECUTION ENGINES (WITH STRICT JSON PARSING)
# ============================================================================
def run_powershell_json(script: str):
    """
    Runs PowerShell cmdlet and parses structured JSON output.
    Guarantees that `parsed` is never a raw string on JSON decode errors,
    preventing AttributeError: 'str' object has no attribute 'get'.
    """
    ps_command = f"$ProgressPreference='SilentlyContinue'; try {{ {script} | ConvertTo-Json -Depth 4 -Compress }} catch {{ Write-Output ('JSON_ERR:' + $_.Exception.Message) }}"
    cmd = [
        "powershell.exe",
        "-NoProfile",
        "-NonInteractive",
        "-ExecutionPolicy", "Bypass",
        "-Command", ps_command
    ]
    try:
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=40)
        stdout = proc.stdout.strip()
        stderr = proc.stderr.strip()

        if proc.returncode != 0 and not stdout:
            return False, None, stdout, stderr or f"Process exited with code {proc.returncode}"

        if not stdout:
            return True, None, "", stderr

        # First direct JSON attempt
        try:
            parsed = json.loads(stdout)
            return True, parsed, stdout, stderr
        except json.JSONDecodeError:
            # Fallback: Find embedded JSON object or array within stdout
            json_match = re.search(r'(\{.*\}|\[.*\])', stdout, re.DOTALL)
            if json_match:
                try:
                    parsed = json.loads(json_match.group(1))
                    return True, parsed, stdout, stderr
                except Exception:
                    pass

            # Do NOT return raw string as parsed JSON - return False with diagnostic
            return False, None, stdout, f"Output not in valid JSON format: {stdout[:120]}"
    except Exception as exc:
        return False, None, "", str(exc)


def run_powershell_raw(script: str):
    """Runs PowerShell command and returns raw console output."""
    cmd = [
        "powershell.exe",
        "-NoProfile",
        "-NonInteractive",
        "-ExecutionPolicy", "Bypass",
        "-Command", script
    ]
    try:
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=30)
        return proc.returncode == 0, proc.stdout.strip(), proc.stderr.strip()
    except Exception as exc:
        return False, "", str(exc)


# ============================================================================
# LIVE WINDOWS RECONNAISSANCE EVIDENCE COLLECTOR (100% LIVE EXTRACTION)
# ============================================================================
def collect_live_evidence(drop_zone_path: str = None, asset_tag: str = None):
    """
    Performs 100% LIVE passive evidence collection from the Windows host.
    Covers all 20 audit control points with type-safe telemetry parsing.
    """
    collected = {}
    print("\n" + "=" * 70)
    print("  COMMENCING 100% LIVE AUDIT RECONNAISSANCE EXTRACTION")
    print("=" * 70)

    # 1. System Hostname, Manufacturer, & Model
    print("  [01/20] Extracting Hardware Baseline (Win32_ComputerSystem)...")
    ok, data, raw_out, err = run_powershell_json(
        "Get-CimInstance Win32_ComputerSystem -ErrorAction SilentlyContinue | Select-Object Name, Manufacturer, Model, TotalPhysicalMemory, Domain"
    )
    if ok and isinstance(data, (dict, list)):
        cs = data if isinstance(data, dict) else (data[0] if data and isinstance(data[0], dict) else {})
        name = cs.get("Name", "Unknown")
        mfg = cs.get("Manufacturer", "Unknown")
        model = cs.get("Model", "Unknown")
        domain = cs.get("Domain", "WORKGROUP")
        mem_bytes = cs.get("TotalPhysicalMemory", 0) or 0
        mem_gb = round(float(mem_bytes) / (1024 ** 3), 2) if mem_bytes else 0

        note = f"Host identified as '{name}' ({mfg} - {model}) in domain/workgroup '{domain}' with {mem_gb} GB RAM."
        if asset_tag:
            note += f" Candidate asset tag for review: '{asset_tag}'."

        collected["HW_SYSTEM"] = {
            "intake_status": "Captured (Ready for Review)",
            "telemetry_summary": f"Hostname: {name} | Manufacturer: {mfg} | Model: {model} | RAM: {mem_gb} GB | Domain: {domain}",
            "recon_notes": note,
            "raw_data": cs,
            "error": err
        }
    else:
        collected["HW_SYSTEM"] = {
            "intake_status": "Requires Reviewer Action",
            "telemetry_summary": "Query on Win32_ComputerSystem returned limited data",
            "recon_notes": f"Diagnostic: {err or raw_out[:120]}",
            "raw_data": raw_out,
            "error": err
        }

    # 2. Total RAM & Disk Capacity / Health
    print("  [02/20] Extracting Storage Volume Headroom (Win32_LogicalDisk)...")
    ok, disks, raw_out, err = run_powershell_json(
        "Get-CimInstance Win32_LogicalDisk -ErrorAction SilentlyContinue | Select-Object DeviceId, DriveType, Size, FreeSpace, FileSystem, VolumeName"
    )
    if ok and isinstance(disks, (dict, list)):
        if isinstance(disks, dict):
            disks = [disks]
        disk_summaries = []
        c_drive_note = ""

        for d in disks:
            if not isinstance(d, dict):
                continue
            dev_id = d.get("DeviceId", "")
            size = d.get("Size") or 0
            free = d.get("FreeSpace") or 0
            fs = d.get("FileSystem", "NTFS")
            vol = d.get("VolumeName", "")
            size_gb = round(float(size) / (1024 ** 3), 1) if size else 0
            free_gb = round(float(free) / (1024 ** 3), 1) if free else 0
            pct_free = round((free_gb / size_gb * 100), 1) if size_gb > 0 else 0

            summary_str = f"{dev_id} [{vol or 'Unlabeled'}] ({fs}): {free_gb} GB free / {size_gb} GB total ({pct_free}% free)"
            disk_summaries.append(summary_str)

            if dev_id.upper().startswith("C"):
                c_drive_note = f"System Drive C: has {free_gb} GB available out of {size_gb} GB ({pct_free}% free headroom)."

        note = c_drive_note if c_drive_note else f"Storage volumes captured: {'; '.join(disk_summaries)}."

        collected["HW_DISK_RAM"] = {
            "intake_status": "Captured (Ready for Review)",
            "telemetry_summary": "; ".join(disk_summaries) if disk_summaries else "No logical volumes returned",
            "recon_notes": note,
            "raw_data": disks,
            "error": err
        }
    else:
        collected["HW_DISK_RAM"] = {
            "intake_status": "Requires Reviewer Action",
            "telemetry_summary": "Storage enumeration returned empty",
            "recon_notes": f"Diagnostic: {err or raw_out[:120]}",
            "raw_data": raw_out,
            "error": err
        }

    # 3. Windows Version & Build Number
    print("  [03/20] Extracting OS Version & Build (Win32_OperatingSystem)...")
    ok, os_info, raw_out, err = run_powershell_json(
        "Get-CimInstance Win32_OperatingSystem -ErrorAction SilentlyContinue | Select-Object Caption, Version, BuildNumber, OSArchitecture, InstallDate"
    )
    if ok and isinstance(os_info, (dict, list)):
        if isinstance(os_info, list) and os_info:
            os_info = os_info[0]
        if isinstance(os_info, dict):
            caption = os_info.get("Caption", "Windows")
            ver = os_info.get("Version", "")
            build = os_info.get("BuildNumber", "")
            arch = os_info.get("OSArchitecture", "")

            note = f"Host running {caption} (Build {build}, {arch}). Ready for review against vendor supported lifecycle list."

            collected["OS_VERSION"] = {
                "intake_status": "Captured (Ready for Review)",
                "telemetry_summary": f"{caption} (Version {ver}, Build {build}, {arch})",
                "recon_notes": note,
                "raw_data": os_info,
                "error": err
            }
        else:
            collected["OS_VERSION"] = {
                "intake_status": "Requires Reviewer Action",
                "telemetry_summary": "OS build metadata non-dictionary format",
                "recon_notes": f"Raw: {raw_out[:120]}",
                "raw_data": raw_out,
                "error": err
            }
    else:
        collected["OS_VERSION"] = {
            "intake_status": "Requires Reviewer Action",
            "telemetry_summary": "OS build metadata uncollected",
            "recon_notes": f"Query failure: {err or raw_out[:120]}",
            "raw_data": raw_out,
            "error": err
        }

    # 4. Last Quality Update Installation Date
    print("  [04/20] Extracting Hotfix & Patch History (Get-HotFix)...")
    ok, hf, raw_out, err = run_powershell_json(
        "Get-HotFix -ErrorAction SilentlyContinue | Sort-Object InstalledOn -Descending | Select-Object -First 10 HotFixID, Description, InstalledOn, InstalledBy"
    )
    if ok and isinstance(hf, (dict, list)):
        if isinstance(hf, dict):
            hf = [hf]
        latest = hf[0] if (hf and isinstance(hf[0], dict)) else {}
        kbid = latest.get("HotFixID", "KB-Unknown")
        installed_on = latest.get("InstalledOn", "")
        desc = latest.get("Description", "Update")

        if isinstance(installed_on, dict) and "DateTime" in installed_on:
            installed_str = installed_on["DateTime"]
        else:
            installed_str = str(installed_on)

        note = f"Discovered {len(hf)} recent hotfixes. Most recent captured patch is {kbid} ({desc}) applied on {installed_str}."

        collected["OS_HOTFIX"] = {
            "intake_status": "Captured (Ready for Review)",
            "telemetry_summary": f"Latest: {kbid} ({desc}) on {installed_str} | Total hotfixes inspected: {len(hf)}",
            "recon_notes": note,
            "raw_data": hf,
            "error": err
        }
    else:
        collected["OS_HOTFIX"] = {
            "intake_status": "Requires Reviewer Action",
            "telemetry_summary": "No patch history returned by Get-HotFix",
            "recon_notes": f"Hotfix query requires elevation or returned no records: {err or raw_out[:120]}",
            "raw_data": raw_out,
            "error": err
        }

    # 5. AV Engine & Definition Status
    print("  [05/20] Extracting Antivirus Engine & Definition Status...")
    ok_av, av_list, _, err_av = run_powershell_json(
        "Get-CimInstance -Namespace root/SecurityCenter2 -ClassName AntivirusProduct -ErrorAction SilentlyContinue | Select-Object displayName, productState, pathToSignedProductExe"
    )
    ok_def, def_status, _, err_def = run_powershell_json(
        "Get-MpComputerStatus -ErrorAction SilentlyContinue | Select-Object AMServiceEnabled, AntispywareEnabled, AntivirusEnabled, AntivirusSignatureLastUpdated, AntivirusSignatureVersion, RealTimeProtectionEnabled"
    )

    av_names = []
    if ok_av and isinstance(av_list, (dict, list)):
        if isinstance(av_list, dict):
            av_list = [av_list]
        for a in av_list:
            if isinstance(a, dict) and a.get("displayName"):
                av_names.append(a.get("displayName"))

    if ok_def and isinstance(def_status, (dict, list)):
        if isinstance(def_status, list) and def_status:
            def_status = def_status[0]
        if isinstance(def_status, dict):
            rtp = def_status.get("RealTimeProtectionEnabled", False)
            sig_ver = def_status.get("AntivirusSignatureVersion", "")
            sig_update = def_status.get("AntivirusSignatureLastUpdated", "")
            if isinstance(sig_update, dict) and "DateTime" in sig_update:
                sig_date_str = sig_update["DateTime"]
            else:
                sig_date_str = str(sig_update)

            note = f"Microsoft Defender active (RealTimeProtection={rtp}, Signature Version {sig_ver}). Definitions updated: {sig_date_str}."
            summary = f"Microsoft Defender (RTP: {rtp}, Sig: {sig_ver}, Updated: {sig_date_str})"
        else:
            summary = "Defender status non-dict format"
            note = "Defender queried."
    elif av_names:
        note = f"Registered SecurityCenter2 Antivirus Provider: {', '.join(av_names)}."
        summary = f"Active AV Provider: {', '.join(av_names)}"
    else:
        note = f"Antivirus state query pending elevated execution. Diagnostic: {err_av or err_def}"
        summary = "Antivirus state pending elevated check"

    collected["SEC_ANTIVIRUS"] = {
        "intake_status": "Captured (Ready for Review)",
        "telemetry_summary": summary,
        "recon_notes": note,
        "raw_data": {"av_products": av_list, "defender_status": def_status},
        "error": err_av or err_def
    }

    # 6. Windows Firewall Profile State
    print("  [06/20] Extracting Windows Firewall Profiles...")
    ok, fw_profiles, raw_out, err = run_powershell_json(
        "Get-NetFirewallProfile -ErrorAction SilentlyContinue | Select-Object Name, Enabled, DefaultInboundAction, DefaultOutboundAction"
    )
    if ok and isinstance(fw_profiles, (dict, list)):
        if isinstance(fw_profiles, dict):
            fw_profiles = [fw_profiles]
        profile_notes = []
        for p in fw_profiles:
            if isinstance(p, dict):
                name = p.get("Name", "")
                en = p.get("Enabled", False)
                in_act = p.get("DefaultInboundAction", "")
                out_act = p.get("DefaultOutboundAction", "")
                profile_notes.append(f"{name}: {'Enabled' if en else 'Disabled'} [In: {in_act}, Out: {out_act}]")

        note = f"Firewall profiles captured: {'; '.join(profile_notes)}."

        collected["SEC_FIREWALL"] = {
            "intake_status": "Captured (Ready for Review)",
            "telemetry_summary": "; ".join(profile_notes) if profile_notes else "No firewall profiles returned",
            "recon_notes": note,
            "raw_data": fw_profiles,
            "error": err
        }
    else:
        collected["SEC_FIREWALL"] = {
            "intake_status": "Requires Reviewer Action",
            "telemetry_summary": "Firewall profiles uncollected",
            "recon_notes": f"Diagnostic: {err or raw_out[:120]}",
            "raw_data": raw_out,
            "error": err
        }

    # 7. Local Users & Enabled Status
    print("  [07/20] Extracting Local User Accounts...")
    ok, users, raw_out, err = run_powershell_json(
        "Get-LocalUser -ErrorAction SilentlyContinue | Select-Object Name, Enabled, LastLogon, Description, PasswordRequired"
    )
    if ok and isinstance(users, (dict, list)):
        if isinstance(users, dict):
            users = [users]
        active_users = []
        disabled_users = []
        for u in users:
            if isinstance(u, dict):
                uname = u.get("Name", "Unknown")
                if u.get("Enabled"):
                    active_users.append(uname)
                else:
                    disabled_users.append(uname)

        note = f"Discovered {len(users)} accounts: {len(active_users)} active ({', '.join(active_users)}), {len(disabled_users)} disabled ({', '.join(disabled_users)})."

        collected["ACC_LOCAL_USERS"] = {
            "intake_status": "Captured (Ready for Review)",
            "telemetry_summary": f"Active: {len(active_users)} ({', '.join(active_users)}) | Disabled: {len(disabled_users)}",
            "recon_notes": note,
            "raw_data": users,
            "error": err
        }
    else:
        ok_net, net_out, net_err = run_powershell_raw("net user")
        collected["ACC_LOCAL_USERS"] = {
            "intake_status": "Captured (Ready for Review)" if ok_net else "Requires Reviewer Action",
            "telemetry_summary": "Captured via fallback 'net user'" if ok_net else "User enumeration failed",
            "recon_notes": f"Captured user list via fallback: {net_out[:120]}..." if ok_net else f"Error: {err or net_err}",
            "raw_data": net_out if ok_net else None,
            "error": err or net_err
        }

    # 8. Local Administrator Group Members
    print("  [08/20] Extracting Local Administrator Group Members...")
    ok, admin_members, raw_out, err = run_powershell_json(
        "Get-LocalGroupMember -Group 'Administrators' -ErrorAction SilentlyContinue | Select-Object Name, PrincipalSource, ObjectClass"
    )
    if ok and isinstance(admin_members, (dict, list)):
        if isinstance(admin_members, dict):
            admin_members = [admin_members]
        names = []
        for m in admin_members:
            if isinstance(m, dict):
                names.append(m.get("Name", "Unknown"))
        note = f"Captured {len(names)} administrator accounts: {', '.join(names)}."

        collected["ACC_ADMIN_MEMBERS"] = {
            "intake_status": "Captured (Ready for Review)",
            "telemetry_summary": f"Administrators ({len(names)}): {', '.join(names)}",
            "recon_notes": note,
            "raw_data": admin_members,
            "error": err
        }
    else:
        ok_lg, lg_out, lg_err = run_powershell_raw("net localgroup Administrators")
        collected["ACC_ADMIN_MEMBERS"] = {
            "intake_status": "Captured (Ready for Review)" if ok_lg else "Requires Reviewer Action",
            "telemetry_summary": "Captured via net localgroup fallback" if ok_lg else "Admin group enumeration failed",
            "recon_notes": "Administrator accounts captured via net localgroup fallback." if ok_lg else f"Error: {err or lg_err}",
            "raw_data": lg_out if ok_lg else None,
            "error": err or lg_err
        }

    # 9. Mapped Network Drives & File Shares
    print("  [09/20] Extracting Mapped Drives & SMB Shares...")
    ok_map, mappings, _, err_map = run_powershell_json(
        "Get-SmbMapping -ErrorAction SilentlyContinue | Select-Object LocalPath, RemotePath, Status"
    )
    ok_shr, shares, _, err_shr = run_powershell_json(
        "Get-SmbShare -ErrorAction SilentlyContinue | Select-Object Name, Path, Description, ShareState"
    )

    map_list = []
    if ok_map and isinstance(mappings, (dict, list)):
        if isinstance(mappings, dict):
            mappings = [mappings]
        for m in mappings:
            if isinstance(m, dict):
                map_list.append(f"{m.get('LocalPath', '?')} -> {m.get('RemotePath', '?')}")

    shr_list = []
    if ok_shr and isinstance(shares, (dict, list)):
        if isinstance(shares, dict):
            shares = [shares]
        for s in shares:
            if isinstance(s, dict) and not str(s.get('Name', '')).endswith('$'):
                shr_list.append(f"{s.get('Name')} ({s.get('Path')})")

    note = ""
    if map_list:
        note += f"Mapped SMB Drives: {', '.join(map_list)}. "
    else:
        note += "No active mapped network drives found. "

    if shr_list:
        note += f"Local non-admin shares: {', '.join(shr_list)}."
    else:
        note += "No custom local network file shares exposed."

    collected["STO_SMB_SHARES"] = {
        "intake_status": "Captured (Ready for Review)",
        "telemetry_summary": f"Mapped Drives: {len(map_list)} | Non-Admin Shares: {len(shr_list)}",
        "recon_notes": note.strip(),
        "raw_data": {"mappings": mappings, "shares": shares},
        "error": err_map or err_shr
    }

    # 10. IP Configuration & Subnet Routing
    print("  [10/20] Extracting IPv4 Network Configuration & Routing...")
    ok_ip, ips, _, err_ip = run_powershell_json(
        "Get-NetIPAddress -AddressFamily IPv4 -ErrorAction SilentlyContinue | Where-Object { $_.IPAddress -notlike '127.*' -and $_.IPAddress -notlike '169.254.*' } | Select-Object InterfaceAlias, IPAddress, PrefixLength"
    )
    ok_rt, route, _, err_rt = run_powershell_json(
        "Get-NetRoute -DestinationPrefix '0.0.0.0/0' -ErrorAction SilentlyContinue | Select-Object NextHop, InterfaceAlias"
    )

    ip_entries = []
    if ok_ip and isinstance(ips, (dict, list)):
        if isinstance(ips, dict):
            ips = [ips]
        for item in ips:
            if isinstance(item, dict):
                ip_entries.append(f"{item.get('InterfaceAlias')}: {item.get('IPAddress')}/{item.get('PrefixLength')}")

    gateway = "Not captured"
    if ok_rt and isinstance(route, (dict, list)):
        if isinstance(route, list) and route and isinstance(route[0], dict):
            gateway = route[0].get("NextHop", "Unknown")
        elif isinstance(route, dict):
            gateway = route.get("NextHop", "Unknown")

    summary_ip = "; ".join(ip_entries) if ip_entries else "No active external IPv4 found"
    note = f"Active interfaces: {summary_ip}. Default Gateway: {gateway}."

    collected["NET_IP_CONFIG"] = {
        "intake_status": "Captured (Ready for Review)" if ip_entries else "Requires Reviewer Action",
        "telemetry_summary": f"IP: {summary_ip} | Gateway: {gateway}",
        "recon_notes": note,
        "raw_data": {"ips": ips, "default_route": route},
        "error": err_ip or err_rt
    }

    # 11. Connected Printers & Print Queues
    print("  [11/20] Extracting Connected Printers & Spoolers...")
    ok, printers, raw_out, err = run_powershell_json(
        "Get-Printer -ErrorAction SilentlyContinue | Select-Object Name, DriverName, PortName, Shared, PrinterStatus"
    )
    if ok and isinstance(printers, (dict, list)):
        if isinstance(printers, dict):
            printers = [printers]
        prn_summaries = []
        for p in printers:
            if isinstance(p, dict):
                prn_summaries.append(f"{p.get('Name')} (Port: {p.get('PortName')})")
        note = f"Discovered {len(printers)} print queue(s): {', '.join(prn_summaries[:4])}."

        collected["PRN_QUEUES"] = {
            "intake_status": "Captured (Ready for Review)",
            "telemetry_summary": f"{len(printers)} print queues registered ({', '.join([p.get('Name','') for p in printers[:3] if isinstance(p, dict)])})",
            "recon_notes": note,
            "raw_data": printers,
            "error": err
        }
    else:
        collected["PRN_QUEUES"] = {
            "intake_status": "Captured (Ready for Review)",
            "telemetry_summary": "No printer queues found on host",
            "recon_notes": "No local or networked print queues returned.",
            "raw_data": None,
            "error": err
        }

    # 12. Router / Network Firewall Configurations & Drop Zone Backups
    print("  [12/20] Inspecting Drop Zone Repository for Firewall Backups...")
    dz_files = []
    search_paths = [
        drop_zone_path,
        r"C:\DropZone",
        r"C:\Backups\Network",
        os.path.expandvars(r"%USERPROFILE%\Documents\DropZone")
    ]
    search_paths = [p for p in search_paths if p and os.path.exists(p)]

    for p in search_paths:
        for root, dirs, files in os.walk(p):
            for f in files:
                if f.endswith((".conf", ".cfg", ".backup", ".xml", ".bin")):
                    dz_files.append(os.path.join(root, f))
        if dz_files:
            break

    if dz_files:
        note = f"Discovered backup file: {os.path.basename(dz_files[0])} in repository ({len(dz_files)} config archives found)."
        status = "Captured (Ready for Review)"
        summary = f"Backup found: {os.path.basename(dz_files[0])}"
    else:
        status = "Pending Reviewer Validation"
        note = "No local .conf/.cfg backup files detected in standard drop paths. Review team to verify central gateway export manually."
        summary = "Drop zone backup awaiting review team verification"

    collected["NET_INFRA_BACKUP"] = {
        "intake_status": status,
        "telemetry_summary": summary,
        "recon_notes": note,
        "raw_data": {"drop_zone_searched": search_paths, "files_found": dz_files},
        "error": None
    }

    # 13. BitLocker Drive Encryption Status
    print("  [13/20] Extracting BitLocker Drive Encryption State...")
    ok_bit, bit_vols, _, err_bit = run_powershell_json(
        "Get-BitLockerVolume -ErrorAction SilentlyContinue | Select-Object MountPoint, VolumeStatus, ProtectionStatus, EncryptionMethod"
    )
    if not ok_bit or not isinstance(bit_vols, (dict, list)):
        ok_bde, bde_out, bde_err = run_powershell_raw("manage-bde -status")
        if ok_bde and "Conversion Status" in bde_out:
            collected["SEC_BITLOCKER"] = {
                "intake_status": "Captured (Ready for Review)",
                "telemetry_summary": "BitLocker status captured via manage-bde",
                "recon_notes": f"Drive encryption telemetry: {bde_out[:140]}...",
                "raw_data": bde_out,
                "error": None
            }
        else:
            collected["SEC_BITLOCKER"] = {
                "intake_status": "Pending Reviewer Validation",
                "telemetry_summary": "BitLocker status unqueried (Requires Administrator)",
                "recon_notes": f"BitLocker query requires elevated privileges: {err_bit or bde_err}",
                "raw_data": None,
                "error": err_bit or bde_err
            }
    else:
        if isinstance(bit_vols, dict):
            bit_vols = [bit_vols]
        bit_summaries = []
        for v in bit_vols:
            if isinstance(v, dict):
                mp = v.get("MountPoint", "?")
                ps = v.get("ProtectionStatus", "Unknown")
                em = v.get("EncryptionMethod", "Unknown")
                vs = v.get("VolumeStatus", "Unknown")
                bit_summaries.append(f"{mp} (Protection: {ps}, Status: {vs}, Method: {em})")

        collected["SEC_BITLOCKER"] = {
            "intake_status": "Captured (Ready for Review)",
            "telemetry_summary": "; ".join(bit_summaries) if bit_summaries else "No BitLocker volumes returned",
            "recon_notes": f"BitLocker volumes telemetry: {'; '.join(bit_summaries)}.",
            "raw_data": bit_vols,
            "error": err_bit
        }

    # 14. NTP Time Synchronization Status
    print("  [14/20] Extracting Time Synchronization & NTP Status...")
    ok_tm, tm_out, tm_err = run_powershell_raw("w32tm /query /status")
    if ok_tm and tm_out and "Source:" in tm_out:
        src = "Unknown"
        stratum = "Unknown"
        for line in tm_out.splitlines():
            if "Source:" in line:
                src = line.split("Source:")[-1].strip()
            if "Stratum:" in line:
                stratum = line.split("Stratum:")[-1].strip()

        collected["OS_TIME_SYNC"] = {
            "intake_status": "Captured (Ready for Review)",
            "telemetry_summary": f"NTP Source: {src} | Stratum: {stratum}",
            "recon_notes": f"System clock synchronized with time source '{src}' (Stratum {stratum}).",
            "raw_data": tm_out,
            "error": None
        }
    else:
        collected["OS_TIME_SYNC"] = {
            "intake_status": "Requires Reviewer Action",
            "telemetry_summary": "w32tm status query returned non-standard output",
            "recon_notes": f"Time service query output: {tm_err or tm_out[:120] or 'Service stopped'}",
            "raw_data": tm_out,
            "error": tm_err
        }

    # 15. User Account Control (UAC) Elevation Policy
    print("  [15/20] Extracting User Account Control (UAC) Registry Policy...")
    ok_uac, uac_data, raw_out, err_uac = run_powershell_json(
        "Get-ItemProperty HKLM:\\Software\\Microsoft\\Windows\\CurrentVersion\\Policies\\System -ErrorAction SilentlyContinue | Select-Object EnableLUA, ConsentPromptBehaviorAdmin"
    )
    if ok_uac and isinstance(uac_data, dict):
        lua = uac_data.get("EnableLUA", "Unknown")
        consent = uac_data.get("ConsentPromptBehaviorAdmin", "Unknown")
        uac_enabled = (lua == 1)
        note = f"UAC EnableLUA={lua} (Enabled: {uac_enabled}), ConsentPromptBehaviorAdmin={consent}."

        collected["SEC_UAC"] = {
            "intake_status": "Captured (Ready for Review)",
            "telemetry_summary": f"UAC Active: {uac_enabled} (EnableLUA={lua}, PromptBehavior={consent})",
            "recon_notes": note,
            "raw_data": uac_data,
            "error": err_uac
        }
    else:
        collected["SEC_UAC"] = {
            "intake_status": "Requires Reviewer Action",
            "telemetry_summary": "UAC policy uncollected",
            "recon_notes": f"Diagnostic: {err_uac or raw_out[:120]}",
            "raw_data": raw_out,
            "error": err_uac
        }

    # 16. Windows Update Service (wuauserv) Health
    print("  [16/20] Extracting Windows Update Service Status...")
    ok_wu, wu_data, raw_out, err_wu = run_powershell_json(
        "Get-Service wuauserv -ErrorAction SilentlyContinue | Select-Object Name, DisplayName, Status, StartType"
    )
    if ok_wu and isinstance(wu_data, dict):
        status_wu = wu_data.get("Status", "Unknown")
        stype = wu_data.get("StartType", "Unknown")
        note = f"Windows Update Service (wuauserv): Status={status_wu}, StartupType={stype}."

        collected["OS_UPDATE_SERVICE"] = {
            "intake_status": "Captured (Ready for Review)",
            "telemetry_summary": f"wuauserv Status: {status_wu} | Startup: {stype}",
            "recon_notes": note,
            "raw_data": wu_data,
            "error": err_wu
        }
    else:
        collected["OS_UPDATE_SERVICE"] = {
            "intake_status": "Requires Reviewer Action",
            "telemetry_summary": "wuauserv service check failed",
            "recon_notes": f"Diagnostic: {err_wu or raw_out[:120]}",
            "raw_data": raw_out,
            "error": err_wu
        }

    # 17. Local Password Policy & Lockout Threshold
    print("  [17/20] Extracting Account Password & Lockout Policy...")
    ok_acc, acc_out, acc_err = run_powershell_raw("net accounts")
    if ok_acc and acc_out and "Minimum password length" in acc_out:
        policy_items = {}
        for line in acc_out.splitlines():
            if ":" in line:
                k, v = line.split(":", 1)
                policy_items[k.strip()] = v.strip()

        min_len = policy_items.get("Minimum password length", "Not specified")
        max_age = policy_items.get("Maximum password age (days)", "Not specified")
        lockout = policy_items.get("Lockout threshold", "Never")

        note = f"Password Policy: Min Length={min_len}, Max Age={max_age} days, Lockout Threshold={lockout}."

        collected["ACC_PASSWORD_POLICY"] = {
            "intake_status": "Captured (Ready for Review)",
            "telemetry_summary": f"Min Length: {min_len} | Max Age: {max_age}d | Lockout Threshold: {lockout}",
            "recon_notes": note,
            "raw_data": policy_items,
            "error": None
        }
    else:
        collected["ACC_PASSWORD_POLICY"] = {
            "intake_status": "Requires Reviewer Action",
            "telemetry_summary": "net accounts execution failed",
            "recon_notes": f"Error querying local password policy: {acc_err or acc_out[:120]}",
            "raw_data": acc_out,
            "error": acc_err
        }

    # 18. Active Listening Network Ports & Exposed Services
    print("  [18/20] Extracting Active Listening TCP Ports...")
    ok_ports, ports_data, raw_out, err_ports = run_powershell_json(
        "Get-NetTCPConnection -State Listen -ErrorAction SilentlyContinue | Select-Object LocalAddress, LocalPort, OwningProcess | Group-Object LocalPort | Select-Object Name, Count"
    )
    if ok_ports and isinstance(ports_data, (dict, list)):
        if isinstance(ports_data, dict):
            ports_data = [ports_data]
        port_nums = []
        for p in ports_data:
            if isinstance(p, dict) and p.get("Name"):
                port_nums.append(str(p.get("Name")))
        note = f"Discovered {len(port_nums)} listening TCP port(s): {', '.join(port_nums[:15])}."

        collected["NET_LISTENING_PORTS"] = {
            "intake_status": "Captured (Ready for Review)",
            "telemetry_summary": f"{len(port_nums)} Listening Ports: {', '.join(port_nums[:10])}",
            "recon_notes": note,
            "raw_data": ports_data,
            "error": err_ports
        }
    else:
        ok_ns, ns_out, ns_err = run_powershell_raw("netstat -ano")
        listening_lines = [l for l in ns_out.splitlines() if "LISTENING" in l.upper()] if ok_ns else []
        collected["NET_LISTENING_PORTS"] = {
            "intake_status": "Captured (Ready for Review)" if ok_ns else "Requires Reviewer Action",
            "telemetry_summary": f"Captured {len(listening_lines)} listening sockets via netstat" if ok_ns else "Listening ports query failed",
            "recon_notes": f"Listening sockets telemetry captured ({len(listening_lines)} active listeners)." if ok_ns else f"Error: {ns_err or raw_out[:120]}",
            "raw_data": ns_out if ok_ns else raw_out,
            "error": err_ports or ns_err
        }

    # 19. Remote Desktop (RDP) & Network Level Authentication (TYPE-SAFE FIX)
    print("  [19/20] Extracting Remote Desktop (RDP) Security Configuration...")
    # Wrap in single PSCustomObject to avoid multi-command semicolon JSON parsing issues
    rdp_ps_script = """
    [PSCustomObject]@{
        fDenyTSConnections = (Get-ItemProperty 'HKLM:\\System\\CurrentControlSet\\Control\\Terminal Server' -Name fDenyTSConnections -ErrorAction SilentlyContinue).fDenyTSConnections
        UserAuthentication = (Get-ItemProperty 'HKLM:\\System\\CurrentControlSet\\Control\\Terminal Server\\WinStations\\RDP-Tcp' -Name UserAuthentication -ErrorAction SilentlyContinue).UserAuthentication
    }
    """
    ok_rdp, rdp_data, raw_out, err_rdp = run_powershell_json(rdp_ps_script)

    if ok_rdp and isinstance(rdp_data, (dict, list)):
        if isinstance(rdp_data, list) and rdp_data:
            rdp_data = rdp_data[0]
        if isinstance(rdp_data, dict):
            f_deny = rdp_data.get("fDenyTSConnections", "Unknown")
            nla = rdp_data.get("UserAuthentication", "Unknown")
            rdp_enabled = (f_deny == 0) if isinstance(f_deny, int) else (f_deny != "Unknown")
            nla_enforced = (nla == 1) if isinstance(nla, int) else False

            note = f"RDP Enabled: {rdp_enabled} (fDenyTSConnections={f_deny}), NLA Enforced: {nla_enforced} (UserAuthentication={nla})."

            collected["SEC_RDP_CONFIG"] = {
                "intake_status": "Captured (Ready for Review)",
                "telemetry_summary": f"RDP Enabled: {rdp_enabled} | NLA Enforced: {nla_enforced}",
                "recon_notes": note,
                "raw_data": rdp_data,
                "error": err_rdp
            }
        else:
            collected["SEC_RDP_CONFIG"] = {
                "intake_status": "Pending Reviewer Validation",
                "telemetry_summary": "RDP registry values retrieved in non-standard format",
                "recon_notes": f"Raw: {raw_out[:120]}",
                "raw_data": raw_out,
                "error": err_rdp
            }
    else:
        # Fallback raw query via reg query
        ok_reg, reg_out, reg_err = run_powershell_raw(
            'reg query "HKLM\\System\\CurrentControlSet\\Control\\Terminal Server" /v fDenyTSConnections'
        )
        if ok_reg and "fDenyTSConnections" in reg_out:
            collected["SEC_RDP_CONFIG"] = {
                "intake_status": "Captured (Ready for Review)",
                "telemetry_summary": "RDP configuration captured via reg query fallback",
                "recon_notes": f"Registry dump: {reg_out.strip()}",
                "raw_data": reg_out,
                "error": None
            }
        else:
            collected["SEC_RDP_CONFIG"] = {
                "intake_status": "Requires Reviewer Action",
                "telemetry_summary": "RDP registry configuration uncollected",
                "recon_notes": f"Diagnostic: {err_rdp or reg_err or raw_out[:120]}",
                "raw_data": raw_out,
                "error": err_rdp or reg_err
            }

    # 20. Security & System Event Log Retention Health
    print("  [20/20] Extracting Security & System Event Log Telemetry...")
    ok_log, log_data, raw_out, err_log = run_powershell_json(
        "Get-WinEvent -ListLog Security, System -ErrorAction SilentlyContinue | Select-Object LogName, RecordCount, IsLogFull, MaximumSizeInBytes"
    )
    if ok_log and isinstance(log_data, (dict, list)):
        if isinstance(log_data, dict):
            log_data = [log_data]
        log_summaries = []
        for l in log_data:
            if isinstance(l, dict):
                lname = l.get("LogName", "")
                rc = l.get("RecordCount", 0)
                isfull = l.get("IsLogFull", False)
                max_mb = round(float(l.get("MaximumSizeInBytes", 0) or 0) / (1024 ** 2), 1)
                log_summaries.append(f"{lname}: {rc} records, MaxSize: {max_mb} MB (Full: {isfull})")

        note = f"Event Log Status: {'; '.join(log_summaries)}."

        collected["LOG_AUDIT_HEALTH"] = {
            "intake_status": "Captured (Ready for Review)",
            "telemetry_summary": "; ".join(log_summaries) if log_summaries else "No log channels returned",
            "recon_notes": note,
            "raw_data": log_data,
            "error": err_log
        }
    else:
        collected["LOG_AUDIT_HEALTH"] = {
            "intake_status": "Requires Reviewer Action",
            "telemetry_summary": "Event log query failed (Requires elevation)",
            "recon_notes": f"Diagnostic: {err_log or raw_out[:120]}",
            "raw_data": raw_out,
            "error": err_log
        }

    print("=" * 70)
    print("  [+] COMPLETED LIVE EXTRACTION FOR ALL 20 AUDIT CONTROLS")
    print("=" * 70 + "\n")
    return collected


# ============================================================================
# EXPORTER 1: PLAIN TEXT RECONNAISSANCE DOSSIER (.TXT)
# ============================================================================
def generate_text_report(audit_data: dict, metadata: dict, output_path: str):
    """Generates an organized, comprehensive text reconnaissance dossier."""
    lines = []
    w = 92

    lines.append("=" * w)
    lines.append("ISO 9001 / ITGC TECHNICAL AUDIT - RECONNAISSANCE INTAKE DOSSIER".center(w))
    lines.append("STAGE 1 PASSIVE TECHNICAL TELEMETRY & BASELINE EVIDENCE INTAKE".center(w))
    lines.append("PREPARED FOR AUDIT REVIEW TEAM EVALUATION (PRELIMINARY DATA)".center(w))
    lines.append("=" * w)
    lines.append(f"Target Hostname   : {metadata['hostname']}")
    lines.append(f"Extraction Mode   : {metadata['mode']}")
    lines.append(f"Extraction Date   : {metadata['timestamp']}")
    lines.append(f"Operating Platform: {metadata['os']}")
    lines.append(f"Auditor / Agent   : {metadata['auditor']}")
    lines.append(f"Compliance Scope  : ISO 9001:2015 (6.1, 7.1.3, 7.1.4, 7.1.5, 7.5.3, 8.5.1, 9.1)")
    lines.append(f"Notice            : Preliminary Reconnaissance Stage. No final pass/fail verdicts.")
    lines.append(f"                    All telemetry packaged with verification prompts for review team.")
    lines.append("-" * w)
    lines.append("")

    # Summary table
    lines.append("RECONNAISSANCE INTAKE CHECKLIST & STATUS MATRIX")
    lines.append("-" * w)
    hdr = f"{'#':<3} | {'ISO Clause':<20} | {'Audit Category':<22} | {'Intake Status':<22} | {'Data Point'}"
    lines.append(hdr)
    lines.append("-" * w)

    ready_count = 0
    pending_count = 0

    for idx, item in enumerate(AUDIT_ITEMS_DEF, start=1):
        item_id = item["id"]
        res = audit_data.get(item_id, {})
        status = res.get("intake_status", "Requires Reviewer Action")

        if "Ready" in status:
            ready_count += 1
        else:
            pending_count += 1

        clause_str = (item["clause"][:18] + "..") if len(item["clause"]) > 20 else item["clause"]
        cat_str = (item["category"][:20] + "..") if len(item["category"]) > 22 else item["category"]
        st_str = (status[:20] + "..") if len(status) > 22 else status
        row = f"{idx:<3} | {clause_str:<20} | {cat_str:<22} | {st_str:<22} | {item['data_point']}"
        lines.append(row)

    lines.append("-" * w)
    lines.append(f"Intake Inventory: {ready_count} Controls Captured & Ready | {pending_count} Items Flagged for Action/Verification")
    lines.append("=" * w)
    lines.append("")

    # Detailed Control Telemetry
    lines.append("DETAILED RECONNAISSANCE TELEMETRY & REVIEWER ACTION ITEMS")
    lines.append("=" * w)

    for idx, item in enumerate(AUDIT_ITEMS_DEF, start=1):
        item_id = item["id"]
        res = audit_data.get(item_id, {})
        status = res.get("intake_status", "Pending")
        note = res.get("recon_notes", "No notes recorded.")
        summary = res.get("telemetry_summary", "N/A")
        reviewer_action = item["reviewer_action"]

        lines.append(f"[{idx:02d}] {item['category'].upper()} :: {item['data_point']}")
        lines.append(f"     Target Clause     : {item['clause']}")
        lines.append(f"     Intake Status     : {status}")
        lines.append(f"     Collection Cmd    : {item['command']}")
        lines.append(f"     Captured Telemetry: {summary}")
        lines.append(f"     Recon Observation : {note}")
        lines.append(f"     REVIEW TEAM ACTION: {reviewer_action}")

        raw_val = res.get("raw_data")
        if raw_val:
            raw_str = json.dumps(raw_val, default=str)
            if len(raw_str) > 160:
                raw_str = raw_str[:157] + "..."
            lines.append(f"     Raw Telemetry Dump: {raw_str}")
        lines.append("-" * w)

    lines.append("")
    lines.append("END OF INTAKE DOSSIER - TRANSMIT TO AUDIT REVIEW TEAM".center(w))
    lines.append("=" * w)

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))

    print(f"[+] Plain text intake dossier generated: {output_path}")
    return output_path


# ============================================================================
# EXPORTER 2: MICROSOFT WORD INTAKE DOSSIER (.DOCX)
# ============================================================================
def set_cell_background(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=70, bottom=70, left=100, right=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def generate_docx_report(audit_data: dict, metadata: dict, output_path: str):
    """Generates an executive Word intake dossier for the review team."""
    if not DOCX_AVAILABLE:
        print("[!] python-docx is not installed. Skipping .docx generation.")
        return None

    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_title = title_p.add_run("ISO 9001 / ITGC Technical Audit Dossier")
    run_title.bold = True
    run_title.font.name = "Arial"
    run_title.font.size = Pt(20)
    run_title.font.color.rgb = RGBColor(16, 44, 87)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(10)
    run_sub = sub_p.add_run("Stage 1 Technical Reconnaissance & Telemetry Preparation for the Audit Review Team")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(10)
    run_sub.font.color.rgb = RGBColor(90, 100, 115)

    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_rows = [
        ("Target Hostname:", f"{metadata['hostname']} (Mode: {metadata['mode']})"),
        ("Collection Timestamp:", metadata['timestamp']),
        ("Audited Platform:", metadata['os']),
        ("Dossier Purpose:", "Reconnaissance Telemetry Intake (Pending Review Team Audit Evaluation)")
    ]

    for i, (k, v) in enumerate(meta_rows):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.8)

        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0)
        set_cell_margins(c1)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(k)
        r0.bold = True
        r0.font.name = "Arial"
        r0.font.size = Pt(8.5)
        r0.font.color.rgb = RGBColor(30, 41, 59)

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(v)
        r1.font.name = "Arial"
        r1.font.size = Pt(8.5)
        r1.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(4)
    r_h2 = h2.add_run("1. Reconnaissance Telemetry Intake Matrix (20 Controls)")
    r_h2.bold = True
    r_h2.font.name = "Arial"
    r_h2.font.size = Pt(12)
    r_h2.font.color.rgb = RGBColor(16, 44, 87)

    sum_table = doc.add_table(rows=1, cols=5)
    sum_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["#", "ISO Clause", "Audit Category", "Target Data Point", "Intake Status"]
    col_widths = [Inches(0.4), Inches(1.4), Inches(1.5), Inches(2.5), Inches(1.2)]

    hdr_row = sum_table.rows[0]
    for idx, heading in enumerate(headers):
        cell = hdr_row.cells[idx]
        cell.width = col_widths[idx]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 80, 80, 70, 70)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(heading)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for i, item in enumerate(AUDIT_ITEMS_DEF, start=1):
        res = audit_data.get(item["id"], {})
        status = res.get("intake_status", "Requires Action")

        row = sum_table.add_row()
        vals = [str(i), item["clause"], item["category"], item["data_point"], status]
        bg_color = "FFFFFF" if i % 2 != 0 else "F8FAFC"

        for c_idx, val in enumerate(vals):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_margins(cell, 50, 50, 70, 70)

            if c_idx == 4:
                if "Ready" in status:
                    set_cell_background(cell, "EFF6FF")
                else:
                    set_cell_background(cell, "FEF3C7")
            else:
                set_cell_background(cell, bg_color)

            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(8)
            if c_idx == 4:
                r.bold = True
                if "Ready" in status:
                    r.font.color.rgb = RGBColor(29, 78, 216)
                else:
                    r.font.color.rgb = RGBColor(180, 83, 9)
            else:
                r.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(4)
    r_h3 = h3.add_run("2. Telemetry Evidence Packages & Reviewer Action Guides")
    r_h3.bold = True
    r_h3.font.name = "Arial"
    r_h3.font.size = Pt(12)
    r_h3.font.color.rgb = RGBColor(16, 44, 87)

    for i, item in enumerate(AUDIT_ITEMS_DEF, start=1):
        res = audit_data.get(item["id"], {})
        status = res.get("intake_status", "Pending")
        note = res.get("recon_notes", "Observation pending.")
        summary = res.get("telemetry_summary", "N/A")
        reviewer_action = item["reviewer_action"]

        p_item = doc.add_paragraph()
        p_item.paragraph_format.space_before = Pt(6)
        p_item.paragraph_format.space_after = Pt(2)
        r_num = p_item.add_run(f"2.{i} {item['category']}: {item['data_point']}")
        r_num.bold = True
        r_num.font.name = "Arial"
        r_num.font.size = Pt(10)
        r_num.font.color.rgb = RGBColor(15, 23, 42)

        card = doc.add_table(rows=5, cols=2)
        card.alignment = WD_TABLE_ALIGNMENT.CENTER
        card_widths = [Inches(1.8), Inches(5.2)]

        details = [
            ("ISO Clause / Scope:", f"{item['clause']} (Status: {status})"),
            ("Collection Command:", f"{item['command']}"),
            ("Captured Telemetry:", f"{summary}"),
            ("Recon Observations:", f"{note}"),
            ("REVIEW TEAM ACTION:", f"{reviewer_action}")
        ]

        for r_i, (k, v) in enumerate(details):
            row = card.rows[r_i]
            c0, c1 = row.cells[0], row.cells[1]
            c0.width = card_widths[0]
            c1.width = card_widths[1]

            if r_i == 4:
                set_cell_background(c0, "FEF3C7")
                set_cell_background(c1, "FFFBEB")
            else:
                set_cell_background(c0, "F1F5F9")
                set_cell_background(c1, "FFFFFF")

            set_cell_margins(c0, 40, 40, 60, 60)
            set_cell_margins(c1, 40, 40, 60, 60)

            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_after = Pt(0)
            r0 = p0.add_run(k)
            r0.bold = True
            r0.font.name = "Arial"
            r0.font.size = Pt(8)
            r0.font.color.rgb = RGBColor(180, 83, 9) if r_i == 4 else RGBColor(71, 85, 105)

            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_after = Pt(0)
            r1 = p1.add_run(v)
            r1.font.name = "Arial"
            r1.font.size = Pt(8)
            if r_i == 4:
                r1.bold = True
                r1.font.color.rgb = RGBColor(146, 64, 14)
            else:
                r1.font.color.rgb = RGBColor(51, 65, 85)

        doc.add_paragraph().paragraph_format.space_after = Pt(3)

    doc.save(output_path)
    print(f"[+] Microsoft Word intake dossier written: {output_path}")
    return output_path


# ============================================================================
# EXPORTER 3: FORMAL RECONNAISSANCE PDF REPORT (.PDF)
# ============================================================================
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 7.5)
        self.setFillColor(colors.HexColor("#64748B"))

        if self._pageNumber > 1:
            self.drawString(54, 750, "ISO 9001 / ITGC Reconnaissance Intake Dossier - Technical Assessment")
            self.setStrokeColor(colors.HexColor("#CBD5E1"))
            self.setLineWidth(0.5)
            self.line(54, 744, 558, 744)

        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(54, 45, 558, 45)
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(558, 32, page_str)
        self.drawString(54, 32, "STAGE 1 TECHNICAL RECONNAISSANCE - FOR AUDIT REVIEW TEAM INTAKE")
        self.restoreState()


def generate_pdf_report(audit_data: dict, metadata: dict, output_path: str):
    """Generates a formal Reconnaissance Intake PDF report covering all 20 controls."""
    if not REPORTLAB_AVAILABLE:
        print("[!] reportlab is not installed. Skipping .pdf generation.")
        return None

    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#0F172A'),
        spaceAfter=3
    )

    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#475569'),
        spaceAfter=8
    )

    h2_style = ParagraphStyle(
        'Heading2_Custom',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#1E3A8A'),
        spaceBefore=6,
        spaceAfter=4
    )

    meta_label = ParagraphStyle(
        'MetaLabel',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor('#334155')
    )

    meta_val = ParagraphStyle(
        'MetaVal',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor('#1E293B')
    )

    th_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7,
        leading=9,
        textColor=colors.white
    )

    tb_style = ParagraphStyle(
        'TableBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7,
        leading=9,
        textColor=colors.HexColor('#1E293B')
    )

    action_label = ParagraphStyle(
        'ActionLabel',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7,
        leading=9,
        textColor=colors.HexColor('#92400E')
    )

    action_val = ParagraphStyle(
        'ActionVal',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7,
        leading=9,
        textColor=colors.HexColor('#78350F')
    )

    story = []

    story.append(Paragraph("ISO 9001 / ITGC Technical Audit: Reconnaissance Dossier", title_style))
    story.append(Paragraph("Stage 1 Passive Evidence Collection & Telemetry Preparation for Audit Review Team", subtitle_style))

    # Metadata Table
    meta_table_data = [
        [
            Paragraph("<b>Target Host:</b>", meta_label), Paragraph(f"{metadata['hostname']} ({metadata['mode']})", meta_val),
            Paragraph("<b>Intake Date:</b>", meta_label), Paragraph(metadata['timestamp'], meta_val)
        ],
        [
            Paragraph("<b>Platform:</b>", meta_label), Paragraph(metadata['os'], meta_val),
            Paragraph("<b>Scope:</b>", meta_label), Paragraph("ISO 9001:2015 & ITGC Technical Baseline (20 Controls)", meta_val)
        ]
    ]
    meta_table = Table(meta_table_data, colWidths=[65, 185, 65, 189])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F1F5F9')),
        ('PADDING', (0, 0), (-1, -1), 3),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LINEBELOW', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 6))

    # Section 1: Summary Matrix
    story.append(Paragraph("1. Reconnaissance Telemetry Intake Matrix (20 Controls)", h2_style))

    matrix_rows = [[
        Paragraph("#", th_style),
        Paragraph("ISO Clause", th_style),
        Paragraph("Audit Category", th_style),
        Paragraph("Required Data Point", th_style),
        Paragraph("Intake Status", th_style)
    ]]

    for idx, item in enumerate(AUDIT_ITEMS_DEF, start=1):
        res = audit_data.get(item["id"], {})
        status = res.get("intake_status", "Requires Action")

        st_color = "#1D4ED8" if "Ready" in status else "#B45309"
        badge_p = Paragraph(f"<font color='{st_color}'><b>{status}</b></font>", tb_style)

        matrix_rows.append([
            Paragraph(str(idx), tb_style),
            Paragraph(item["clause"], tb_style),
            Paragraph(item["category"], tb_style),
            Paragraph(item["data_point"], tb_style),
            badge_p
        ])

    table_matrix = Table(matrix_rows, colWidths=[16, 85, 95, 218, 90])
    ts = TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('PADDING', (0, 0), (-1, -1), 2.5),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
    ])
    for r in range(1, len(matrix_rows)):
        if r % 2 == 0:
            ts.add('BACKGROUND', (0, r), (-2, r), colors.HexColor('#F8FAFC'))
    table_matrix.setStyle(ts)
    story.append(table_matrix)
    story.append(Spacer(1, 8))

    # Section 2: Detailed Evidence Cards & Review Actions
    story.append(Paragraph("2. Telemetry Packages & Reviewer Action Items", h2_style))

    for idx, item in enumerate(AUDIT_ITEMS_DEF, start=1):
        res = audit_data.get(item["id"], {})
        status = res.get("intake_status", "Pending")
        note = res.get("recon_notes", "Observation pending.")
        summary = res.get("telemetry_summary", "N/A")
        reviewer_action = item["reviewer_action"]

        st_color = "#1D4ED8" if "Ready" in status else "#B45309"

        card_data = [
            [
                Paragraph(f"<b>[{idx:02d}] {item['category']} — {item['data_point']}</b>", meta_label),
                Paragraph(f"<font color='{st_color}'><b>{status}</b></font>", meta_label)
            ],
            [
                Paragraph("<b>ISO Clause:</b>", meta_label),
                Paragraph(item['clause'], meta_val)
            ],
            [
                Paragraph("<b>Collection Command:</b>", meta_label),
                Paragraph(f"<code>{item['command']}</code>", tb_style)
            ],
            [
                Paragraph("<b>Captured Telemetry:</b>", meta_label),
                Paragraph(summary, tb_style)
            ],
            [
                Paragraph("<b>Recon Observation:</b>", meta_label),
                Paragraph(f"<i>{note}</i>", tb_style)
            ],
            [
                Paragraph("<b>Reviewer Action:</b>", action_label),
                Paragraph(f"<b>{reviewer_action}</b>", action_val)
            ]
        ]

        card_table = Table(card_data, colWidths=[105, 399])
        card_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#EFF6FF')),
            ('BACKGROUND', (0, 1), (0, -2), colors.HexColor('#F8FAFC')),
            ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#FEF3C7')),
            ('PADDING', (0, 0), (-1, -1), 2.2),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('INNERGRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#E2E8F0')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP')
        ]))

        story.append(KeepTogether([card_table, Spacer(1, 4)]))

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"[+] PDF intake dossier written: {output_path}")
    return output_path


# ============================================================================
# EXPORTER 4: ZIP ARCHIVE BUNDLER
# ============================================================================
def create_audit_zip(generated_files: list, zip_path: str):
    """Bundles all generated intake reports and raw telemetry into a ZIP archive."""
    with zipfile.ZipFile(zip_path, 'w', compression=zipfile.ZIP_DEFLATED) as zf:
        for file_path in generated_files:
            if file_path and os.path.exists(file_path):
                arcname = os.path.basename(file_path)
                zf.write(file_path, arcname=arcname)
                print(f"  -> Added to archive: {arcname}")

    print(f"[+] Consolidated reconnaissance package created: {zip_path} ({os.path.getsize(zip_path)} bytes)")
    return zip_path


# ============================================================================
# MAIN ENTRYPOINT & CLI ORCHESTRATOR
# ============================================================================
def main():
    parser = argparse.ArgumentParser(
        description="ISO 9001 / ITGC Technical Audit - 100% Live Reconnaissance Intake Dossier Generator"
    )
    parser.add_argument(
        "--output-dir",
        default="./audit_reports",
        help="Directory to save generated reconnaissance reports and zip bundle"
    )
    parser.add_argument(
        "--asset-tag",
        default="AGA-AQUA",
        help="Physical hardware asset tag to register for review team cross-referencing"
    )
    parser.add_argument(
        "--drop-zone-path",
        default=None,
        help="Custom filesystem path to search for firewall/router backup configurations (.conf)"
    )

    args = parser.parse_args()

    os.makedirs(args.output_dir, exist_ok=True)
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    prefix = f"ISO9001_Live_Recon_{timestamp}"

    is_windows = platform.system().lower() == "windows"

    if not is_windows:
        print("\n" + "!" * 75)
        print("  WARNING: Non-Windows operating system detected (" + platform.system() + ").")
        print("  This audit script is engineered for 100% LIVE Windows extraction.")
        print("  To perform the live audit extraction, run this script directly on the target")
        print("  Windows host in an Elevated PowerShell / Command Prompt window.")
        print("!" * 75 + "\n")
        hostname = platform.node()
        os_info = f"{platform.system()} {platform.release()} ({platform.machine()})"
        mode_label = "Local System (Non-Windows Diagnostic)"
        audit_data = {}
        for item in AUDIT_ITEMS_DEF:
            audit_data[item["id"]] = {
                "intake_status": "Requires Live Windows Execution",
                "telemetry_summary": f"Host is running {platform.system()}. Run on target Windows machine for live telemetry.",
                "recon_notes": f"This control requires live Windows PowerShell execution ({item['command']}).",
                "raw_data": {"platform": platform.platform(), "node": platform.node()},
                "error": "Non-Windows OS"
            }
    else:
        print("[*] Running on native Windows host. Commencing 100% LIVE passive extraction...")
        audit_data = collect_live_evidence(drop_zone_path=args.drop_zone_path, asset_tag=args.asset_tag)
        hostname = audit_data.get("HW_SYSTEM", {}).get("raw_data", {}).get("Name", platform.node())
        os_info = audit_data.get("OS_VERSION", {}).get("telemetry_summary", platform.platform())
        mode_label = "100% LIVE Windows Audit Telemetry"

    metadata = {
        "hostname": hostname,
        "os": os_info,
        "mode": mode_label,
        "timestamp": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "auditor": "Automated Passive Audit Collector (Stage 1 Intake / Gemini Spark)"
    }

    json_path = os.path.join(args.output_dir, f"{prefix}_raw_telemetry.json")
    with open(json_path, "w", encoding="utf-8") as jf:
        json.dump({"metadata": metadata, "intake_items": audit_data}, jf, indent=2, default=str)

    txt_path = os.path.join(args.output_dir, f"{prefix}.txt")
    docx_path = os.path.join(args.output_dir, f"{prefix}.docx")
    pdf_path = os.path.join(args.output_dir, f"{prefix}.pdf")
    zip_path = os.path.join(args.output_dir, f"{prefix}_package.zip")

    generated_files = [json_path]

    print("\n--- Generating Reconnaissance Reports ---")
    gen_txt = generate_text_report(audit_data, metadata, txt_path)
    if gen_txt:
        generated_files.append(gen_txt)

    gen_docx = generate_docx_report(audit_data, metadata, docx_path)
    if gen_docx:
        generated_files.append(gen_docx)

    gen_pdf = generate_pdf_report(audit_data, metadata, pdf_path)
    if gen_pdf:
        generated_files.append(gen_pdf)

    final_zip = create_audit_zip(generated_files, zip_path)

    print("\n=================================================================")
    print("AUDIT RECONNAISSANCE DOSSIER GENERATION COMPLETE")
    print("=================================================================")
    print(f"1. Plain Text Dossier : {txt_path}")
    print(f"2. Word Intake Doc    : {docx_path if gen_docx else 'N/A'}")
    print(f"3. PDF Intake Report  : {pdf_path if gen_pdf else 'N/A'}")
    print(f"4. Raw JSON Telemetry : {json_path}")
    print(f"5. Final Zip Package  : {final_zip}")
    print("=================================================================\n")


if __name__ == "__main__":
    main()
